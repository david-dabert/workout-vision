#!/usr/bin/env python3
"""
Convert official ANSES CIQUAL 2025 dataset to compact JSON format.

Dataset: Table de composition nutritionnelle des aliments Ciqual 2025 (ANSES)
DOI: 10.57745/RDMHWY
Source files:
  - Table Ciqual 2025_FR_2025_11_03.xlsx (id 666260, SHA256: 5555c572fa3735991298d832d0427788fa69a11b4fd20a5d580d58942369fbb0)
  - alim_2025_11_03.xml (id 666252, SHA256: e0b1de25b3039028205e9d54a96892e403e1b313c2efeb41180fabe132627478)
Licence: Licence Ouverte / Open Licence 2.0 (Etalab)
"""

import hashlib
import json
import os
import sys
import urllib.request
import xml.etree.ElementTree as ET
import openpyxl

XLSX_URL = "https://entrepot.recherche.data.gouv.fr/api/access/datafile/666260"
XML_URL = "https://entrepot.recherche.data.gouv.fr/api/access/datafile/666252"

CACHE_DIR = "/tmp/ciqual_2025_cache"
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_JSON = os.path.join(SCRIPT_DIR, "ciqual-compact.json")


def download_file(url: str, dest_path: str, expected_sha256: str = None) -> bytes:
    if os.path.exists(dest_path) and os.path.getsize(dest_path) > 0:
        with open(dest_path, "rb") as f:
            data = f.read()
        if expected_sha256:
            actual = hashlib.sha256(data).hexdigest()
            if actual == expected_sha256:
                return data
    print(f"Downloading from {url} to {dest_path}...")
    req = urllib.request.Request(url, headers={"User-Agent": "WorkoutVision-CiqualConverter/1.0"})
    import ssl
    try:
        import certifi
        ctx = ssl.create_default_context(cafile=certifi.where())
    except Exception:
        ctx = ssl.create_default_context()
        ctx.check_hostname = False
        ctx.verify_mode = ssl.CERT_NONE
    with urllib.request.urlopen(req, context=ctx) as resp:
        data = resp.read()
    with open(dest_path, "wb") as f:
        f.write(data)
    if expected_sha256:
        actual = hashlib.sha256(data).hexdigest()
        assert actual == expected_sha256, f"SHA256 mismatch for {dest_path}: got {actual}, expected {expected_sha256}"
    return data


def parse_numeric(v):
    if v is None:
        return None, ""
    s = str(v).strip().replace("\n", " ")
    if not s or s == "-" or "traces" in s.lower() or "<" in s:
        return None, s
    try:
        val = float(s.replace(",", "."))
        return val, s
    except ValueError:
        return None, s


def main():
    os.makedirs(CACHE_DIR, exist_ok=True)
    xlsx_path = os.path.join(CACHE_DIR, "Table_Ciqual_2025_FR_2025_11_03.xlsx")
    xml_path = os.path.join(CACHE_DIR, "alim_2025_11_03.xml")

    # Download source files if not cached
    download_file(XLSX_URL, xlsx_path, "5555c572fa3735991298d832d0427788fa69a11b4fd20a5d580d58942369fbb0")
    download_file(XML_URL, xml_path, "e0b1de25b3039028205e9d54a96892e403e1b313c2efeb41180fabe132627478")

    # Parse English names from XML
    print("Parsing English food names from XML...")
    tree = ET.parse(xml_path)
    eng_names = {}
    for elem in tree.getroot().findall("ALIM"):
        code_el = elem.find("alim_code")
        eng_el = elem.find("alim_nom_eng")
        if code_el is not None and code_el.text:
            code_str = code_el.text.strip()
            if eng_el is not None and eng_el.text:
                eng_names[code_str] = eng_el.text.strip()

    print("Opening Excel table...")
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["composition nutritionnelle"]

    foods = []
    # Columns in Excel (1-based):
    # 4: alim_grp_nom_fr
    # 7: alim_code
    # 8: alim_nom_fr
    # 11: Energie, Règlement UE N° 1169/2011 (kcal/100 g)
    # 16: Protéines, N x 6.25 (g/100 g)
    # 17: Glucides (g/100 g)
    # 18: Lipides (g/100 g)
    # 27: Fibres alimentaires (g/100 g)

    for r in range(2, ws.max_row + 1):
        code_val = ws.cell(r, 7).value
        if code_val is None:
            continue
        code = str(code_val).strip()
        name_fr = str(ws.cell(r, 8).value or "").strip().replace("\n", " ")
        group = str(ws.cell(r, 4).value or "").strip().replace("\n", " ")
        name_en = eng_names.get(code)

        e_val, e_raw = parse_numeric(ws.cell(r, 11).value)
        p_val, p_raw = parse_numeric(ws.cell(r, 16).value)
        c_val, c_raw = parse_numeric(ws.cell(r, 17).value)
        fat_val, fat_raw = parse_numeric(ws.cell(r, 18).value)
        fib_val, fib_raw = parse_numeric(ws.cell(r, 27).value)

        entry = {
            "code": code,
            "name_fr": name_fr,
            "name_en": name_en,
            "group": group,
            "energy_kcal": e_val,
            "protein_g": p_val,
            "protein": p_val,
            "carbohydrate_g": c_val,
            "carbohydrate": c_val,
            "fat_g": fat_val,
            "fat": fat_val,
            "fibre_g": fib_val,
            "fibre": fib_val,
            "raw": {
                "energy_kcal": e_raw,
                "protein_g": p_raw,
                "carbohydrate_g": c_raw,
                "fat_g": fat_raw,
                "fibre_g": fib_raw,
            },
        }
        foods.append(entry)

    print(f"Total foods parsed: {len(foods)}")
    with open(OUTPUT_JSON, "w", encoding="utf-8") as f:
        json.dump(foods, f, ensure_ascii=False, indent=2)

    json_size_mb = os.path.getsize(OUTPUT_JSON) / (1024 * 1024)
    print(f"Written {OUTPUT_JSON} ({json_size_mb:.2f} MB)")


if __name__ == "__main__":
    main()
